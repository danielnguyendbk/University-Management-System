import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()

# Sheet 1: HUONG_DAN
ws1 = wb.active
ws1.title = "HUONG_DAN"
header_font = Font(bold=True)
header_fill = PatternFill("solid", fgColor="C5D9F1")

ws1.append(["HƯỚNG DẪN IMPORT LỊCH THI"])
ws1["A1"].font = header_font
ws1["A1"].fill = header_fill
ws1.append(["- Admin KHÔNG nhập sĩ số và dãy ghế. Sĩ số do backend tự tính dựa trên enrollments."])
ws1.append(["- File gồm các sheet: EXAMS_IMPORT (nhập lịch thi), INVIGILATORS_IMPORT (nhập giám thị)."])
ws1.column_dimensions["A"].width = 100

# Sheet 2: EXAMS_IMPORT
ws2 = wb.create_sheet("EXAMS_IMPORT")
exam_headers = ["action", "semester_id", "section_code", "room_code", "exam_type", "exam_method", "exam_date", "start_time", "end_time", "status", "note"]
ws2.append(exam_headers)
for cell in ws2[1]:
    cell.font = header_font
    cell.fill = header_fill

# Dummy data for exams
ws2.append(["ADD", 1, "INT1306_01", "701", "midterm", "WRITTEN", "2026-06-01", "08:00", "09:30", "SCHEDULED", "Thi giữa kỳ mẫu"])
ws2.append(["ADD", 1, "INT1306_02", "702", "midterm", "WRITTEN", "2026-06-01", "08:00", "09:30", "SCHEDULED", "Thi giữa kỳ mẫu"])

for col in ws2.columns:
    ws2.column_dimensions[col[0].column_letter].width = 15

# Sheet 3: INVIGILATORS_IMPORT
ws3 = wb.create_sheet("INVIGILATORS_IMPORT")
inv_headers = ["semester_id", "section_code", "exam_date", "start_time", "room_code", "lecturer_code", "invigilator_role", "note"]
ws3.append(inv_headers)
for cell in ws3[1]:
    cell.font = header_font
    cell.fill = header_fill

# Dummy data for invigilators
ws3.append([1, "INT1306_01", "2026-06-01", "08:00", "701", "GV01", "MAIN", ""])
ws3.append([1, "INT1306_02", "2026-06-01", "08:00", "702", "GV02", "MAIN", ""])

for col in ws3.columns:
    ws3.column_dimensions[col[0].column_letter].width = 15

wb.save("UMS_Import_Exam_Template_Sample.xlsx")
